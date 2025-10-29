# app/api/documentos_url.py
from fastapi import APIRouter, Depends, HTTPException, status
from pydantic import BaseModel, HttpUrl
from sqlalchemy.orm import Session
from app.core.database import get_db
from app.api.auth import get_current_user
from app.models.documento import Documento
from app.services.documentos_url_service import process_external_document
from datetime import datetime
from email.utils import parsedate_to_datetime
import os

router = APIRouter(tags=["Documentos desde URL"])

class URLRequest(BaseModel):
    url: HttpUrl
    version: str | None = "1.0"

@router.post("/desde-url")
def desde_url(req: URLRequest, db: Session = Depends(get_db), usuario = Depends(get_current_user)):
    """
    Recibe una URL pública (Drive/OneDrive/directa), descarga el archivo,
    extrae metadatos y lo registra en la tabla 'documentos'.
    """
    try:
        metadata = process_external_document(str(req.url), usuario.id, req.version or "1.0")
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"No se pudo procesar la URL: {str(e)}")

    # Verificar duplicado por hash
    hash_val = metadata.get("hash_archivo")
    if not hash_val:
        raise HTTPException(status_code=500, detail="No se pudo calcular el hash del archivo")

    existe = db.query(Documento).filter(Documento.hash_archivo == hash_val).first()
    duplicado = bool(existe)

    # Convertir last_modified HTTP a datetime
    last_modified_str = metadata.get("last_modified")
    if last_modified_str:
        try:
            creado_en = parsedate_to_datetime(last_modified_str)
        except Exception:
            creado_en = datetime.utcnow()
    else:
        creado_en = datetime.utcnow()

    # Intentar detectar autor automáticamente
    autor_detectado = None
    try:
        ruta = metadata.get("ruta_guardado", "")
        if os.path.exists(ruta):
            ext = metadata.get("extension", "").lower()

            # Si es un DOCX/DOC: extrae metadatos internos
            if ext in ("docx", "doc"):
                from docx import Document as DocxDocument
                docx_file = DocxDocument(ruta)
                core = docx_file.core_properties
                autor_detectado = core.author or core.last_modified_by

            # Si es un PDF: intenta leer autor desde metadatos del PDF
            elif ext == "pdf":
                from PyPDF2 import PdfReader
                pdf = PdfReader(ruta)
                info = pdf.metadata or {}
                autor_detectado = info.get("/Author") or info.get("Author")

            # Si es un Excel (xlsx o xlsm): toma creador/modificador
            elif ext in ("xlsx", "xlsm"):
                from openpyxl import load_workbook
                wb = load_workbook(ruta, read_only=True)
                props = wb.properties
                autor_detectado = props.creator or props.lastModifiedBy

    except Exception as e:
        print(f"[Advertencia] No se pudo leer metadatos de autor: {e}")

    # Si no se detecta autor, usar el usuario autenticado
    if not autor_detectado:
        autor_detectado = getattr(usuario, "nombre", None) or getattr(usuario, "email", None) or "Desconocido"

    nuevo = Documento(
        nombre_archivo=metadata.get("nombre_archivo", "sin_nombre"),
        extension=metadata.get("extension", ""),
        version=metadata.get("version", "1.0"),
        hash_archivo=hash_val,
        ruta_guardado=metadata.get("ruta_guardado", ""),
        tamano_kb=float(metadata.get("tamano_kb", 0)),
        duplicado=duplicado,
        usuario_id=usuario.id,
        creado_en=creado_en,
        content_type=metadata.get("content_type"),
        last_modified=metadata.get("last_modified"),
        servidor=metadata.get("servidor"),
        tipo_documento=metadata.get("tipo_documento"),
        categoria=metadata.get("categoria"),
        confidencialidad=metadata.get("confidencialidad"),
        autor=autor_detectado
    )

    db.add(nuevo)
    db.commit()
    db.refresh(nuevo)

    # Preparar respuesta con todos los metadatos
    response = {
        "status": "ok",
        "mensaje": "Documento registrado correctamente",
        "documento": {
            "id": nuevo.id,
            "nombre": nuevo.nombre_archivo,
            "extension": nuevo.extension,
            "version": nuevo.version,
            "tamano_kb": nuevo.tamano_kb,
            "ruta_guardado": nuevo.ruta_guardado,
            "duplicado": nuevo.duplicado,
            "creado_en": nuevo.creado_en.isoformat() if nuevo.creado_en else None,
            "autor": nuevo.autor
        },
        "clasificacion": {
            "tipo_documento": nuevo.tipo_documento,
            "categoria": nuevo.categoria,
            "confidencialidad": nuevo.confidencialidad,
            "autor": nuevo.autor
        },
        "metadatos_extra": {
            k: v for k, v in metadata.items() if k not in (
                "nombre_archivo", "extension", "version", "hash_archivo", "ruta_guardado", "tamano_kb",
                "tipo_documento", "categoria", "confidencialidad", "autor"
            )
        }
    }

    return response